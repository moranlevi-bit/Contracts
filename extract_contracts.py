"""Collect key data from contracts into a table.

Reads every .pdf, .docx and .txt file in a folder, asks Claude to extract
Org Name, Industry, Earnings and Bonuses, and writes one row per contract
to an Excel (.xlsx) or CSV file.

Usage:
    python extract_contracts.py contracts/ -o contracts_table.xlsx
"""

import argparse
import base64
import csv
import json
import sys
from pathlib import Path

import anthropic
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font

MODEL = "claude-opus-5-5"

COLUMNS = ["File", "Org Name", "Industry", "Earnings", "Bonuses"]

SCHEMA = {
    "type": "object",
    "properties": {
        "org_name": {
            "type": ["string", "null"],
            "description": "Name of the organization / company that is party to the contract",
        },
        "industry": {
            "type": ["string", "null"],
            "description": "Industry or sector of the organization",
        },
        "earnings": {
            "type": ["string", "null"],
            "description": "Base pay, salary, fees or contract value, with amount, currency and period",
        },
        "bonuses": {
            "type": ["string", "null"],
            "description": "Bonuses, commissions or incentive payments, with amounts and conditions",
        },
    },
    "required": ["org_name", "industry", "earnings", "bonuses"],
    "additionalProperties": False,
}

PROMPT = """Extract the following from the contract above:
- org_name: the organization / company that is party to the contract
- industry: the organization's industry or sector (infer from the contract if not stated)
- earnings: base pay, salary, fees or contract value - include amount, currency and period
- bonuses: bonuses, commissions or incentive payments - include amounts and conditions

Use null for anything the contract does not contain. Do not guess amounts."""

SUPPORTED = {".pdf", ".docx", ".txt"}


def contract_content(path: Path) -> dict:
    """Return the contract as a content block for the request."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        data = base64.standard_b64encode(path.read_bytes()).decode()
        return {
            "type": "document",
            "source": {"type": "base64", "media_type": "application/pdf", "data": data},
        }
    if suffix == ".docx":
        doc = Document(str(path))
        text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                text += "\n" + " | ".join(cell.text for cell in row.cells)
    else:
        text = path.read_text(encoding="utf-8", errors="replace")
    return {"type": "text", "text": f"<contract>\n{text}\n</contract>"}


def extract(client: anthropic.Anthropic, path: Path) -> dict:
    response = client.beta.messages.create(
        model=MODEL,
        max_tokens=16000,
        betas=["server-side-fallback-2026-07-01"],
        fallbacks="default",
        output_config={
            "effort": "medium",
            "format": {"type": "json_schema", "schema": SCHEMA},
        },
        messages=[
            {
                "role": "user",
                "content": [contract_content(path), {"type": "text", "text": PROMPT}],
            }
        ],
    )
    if response.stop_reason == "refusal":
        raise RuntimeError("request was declined")
    if response.stop_reason == "max_tokens":
        raise RuntimeError("response was cut off (max_tokens)")
    text = next(b.text for b in response.content if b.type == "text")
    return json.loads(text)


def write_xlsx(rows: list[list], out: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Contracts"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for col, width in zip("ABCDE", (30, 30, 25, 50, 50)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(out)


def write_csv(rows: list[list], out: Path) -> None:
    with out.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(COLUMNS)
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("folder", type=Path, help="folder containing contracts")
    parser.add_argument(
        "-o", "--output", type=Path, default=Path("contracts_table.xlsx"),
        help="output file, .xlsx or .csv (default: contracts_table.xlsx)",
    )
    args = parser.parse_args()

    files = sorted(p for p in args.folder.rglob("*") if p.suffix.lower() in SUPPORTED)
    if not files:
        sys.exit(f"No .pdf, .docx or .txt files found in {args.folder}")

    client = anthropic.Anthropic()
    rows = []
    for path in files:
        print(f"Extracting {path.name}...", file=sys.stderr)
        try:
            data = extract(client, path)
            rows.append([path.name, data["org_name"], data["industry"],
                         data["earnings"], data["bonuses"]])
        except (anthropic.APIError, RuntimeError, ValueError) as e:
            print(f"  failed: {e}", file=sys.stderr)
            rows.append([path.name, f"ERROR: {e}", None, None, None])

    if args.output.suffix.lower() == ".csv":
        write_csv(rows, args.output)
    else:
        write_xlsx(rows, args.output)
    print(f"Wrote {len(rows)} rows to {args.output}", file=sys.stderr)


if __name__ == "__main__":
    main()
